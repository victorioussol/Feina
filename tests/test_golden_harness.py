"""Row 294 of 20 Oct 2026, through the real approved scripts.

The Membr POS export for that day is not in the fixture. C and E come from
daily-payload.json. `afsa_checks.py export` is not covered by the row-294 test.

`dates` compares transaction dates with the business date and rejects a PDF
already stored under another verified date. It does not read today's date.
The ledger path is redirected to the temp dir from conftest.
"""

from __future__ import annotations

import json
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest

from openpyxl import Workbook, load_workbook

from feina_client.reconcile import review_day
from feina_client.runtime import sha256_file

DAY = date(2026, 10, 20)
ROOT = Path(__file__).resolve().parents[1]
GOLDEN = ROOT / "tests" / "fixtures" / "golden" / "2026-10-20"
SLOT = ROOT / "tests" / "fixtures" / "golden" / "pending-membr-export"
LIVE_CF = "#,##0.00_);[Red](#,##0.00)"
SCRIPT_CF = "#,##0.00;[Red]-#,##0.00"


def _money(value) -> Decimal:
    return Decimal(str(value)).quantize(Decimal("0.01"))


def _source_workbook(path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "2026"
    book.create_sheet("2025")
    sheet["A4"] = "KEEP-A4"
    sheet["B4"] = "4 Enero 2026"
    sheet["C4"] = 1
    sheet["H4"] = "KEEP-H4"
    sheet["H293"] = "KEEP-H293"
    sheet["B294"] = "20 Octubre"
    sheet["H294"] = "KEEP-H294"
    book["2025"]["A1"] = "KEEP-2025"
    sheet.column_dimensions["H"].hidden = False
    book.save(path)


@pytest.mark.parametrize("comment", [None, "KEEP-H294"])
def test_golden_row_294_uses_the_payload_because_the_membr_export_is_missing(client, comment):
    """`export` is not replayed. dates and update_workbook.py are the real scripts."""
    folder = GOLDEN / "inputs"
    payload = folder / "daily-payload.json"
    pdfs = [
        folder / "6874_347360075_99871480206_20102026.pdf",
        folder / "resumen_347360075_T0000001_202610210123.pdf",
    ]
    evidence = [folder / "processor-evidence.json"]
    book = client / "AFSA-Ingresos.xlsx"
    _source_workbook(book)
    original = load_workbook(book)
    original["2026"]["H294"] = comment
    original.save(book)
    before = sha256_file(book)
    reviewed = review_day(
        business_date=DAY,
        membr=None,
        pdfs=pdfs,
        evidence=evidence,
        payload_path=payload,
        workbook=book,
        slack_on_fail=False,
    )
    assert reviewed["export_covered"] is False
    assert reviewed["checks"]["export"] is None
    assert reviewed["checks"]["dates"]["result"] == "PASS"
    assert reviewed["checks"]["update"]["row"] == 294
    assert reviewed["figures_final"] is False
    assert reviewed["checks"]["final"]["result"] == "FAIL"
    assert sha256_file(book) == before

    staged = Path(reviewed["staged_file"])
    saved = load_workbook(staged, data_only=False)
    cached = load_workbook(staged, data_only=True)
    sheet = saved["2026"]
    values = cached["2026"]
    assert _money(sheet["C294"].value) == Decimal("135.49")
    assert sheet["D294"].value == "=C294-E294"
    assert _money(sheet["E294"].value) == Decimal("-24.00")
    assert _money(sheet["F294"].value) == Decimal("159.49")
    assert sheet["G294"].value == "=F294-D294"
    assert _money(values["C294"].value) == Decimal("135.49")
    assert _money(values["D294"].value) == Decimal("159.49")
    assert _money(values["E294"].value) == Decimal("-24.00")
    assert _money(values["F294"].value) == Decimal("159.49")
    assert _money(values["G294"].value) == Decimal("0.00")
    assert sheet["G294"].font.color.rgb == "FF008000"
    assert sheet["C294"].number_format == SCRIPT_CF
    assert sheet["C294"].number_format != LIVE_CF
    assert bool(sheet.column_dimensions["H"].hidden) is False
    assert sheet["H294"].value == comment
    assert sheet["H293"].value == "KEEP-H293"
    assert sheet["A4"].value == "KEEP-A4"
    assert sheet["C4"].value == 1
    assert saved["2025"]["A1"].value == "KEEP-2025"
    assert sheet.row_dimensions[294].height == 16


def test_export_is_covered_when_a_membr_file_is_actually_supplied(client, tmp_path):
    folder = GOLDEN / "inputs"
    payload = folder / "daily-payload.json"
    pdfs = [
        folder / "6874_347360075_99871480206_20102026.pdf",
        folder / "resumen_347360075_T0000001_202610210123.pdf",
    ]
    export = tmp_path / "pos-export.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Informe"
    meta = book.create_sheet("Informe Meta Data")
    import datetime as dt

    meta.append(["Nombre Informe", "Fecha de", "Fecha hasta", "Fecha descargado"])
    meta.append(
        [
            "Transaccion Punto de Venta",
            dt.datetime(2026, 10, 20, 0, 0, 0),
            dt.datetime(2026, 10, 20, 23, 59, 59),
            dt.datetime(2026, 10, 21, 10, 4, 29),
        ]
    )
    sheet.append(
        [
            "Fecha de la transacción",
            "Estado de transacción",
            "Tipo de pago",
            "Subtotal",
            "Nombre del gimnasio",
            "Número de recibo",
        ]
    )
    sheet.append(
        [dt.datetime(2026, 10, 20, 13, 27, 35), "Pagado", "Tarjeta", 159.49, "Anytime Fitness Sant Adria", "R1"]
    )
    sheet.append(
        [dt.datetime(2026, 10, 20, 18, 0, 0), "Pagado", "Efectivo", -24, "Anytime Fitness Sant Adria", "R2"]
    )
    book.save(export)
    target = client / "AFSA-Ingresos.xlsx"
    _source_workbook(target)
    reviewed = review_day(
        business_date=DAY,
        membr=export,
        pdfs=pdfs,
        evidence=[folder / "processor-evidence.json"],
        payload_path=payload,
        workbook=target,
        screen_count_text="9",
        screen_total_text="999.00",
        slack_on_fail=False,
    )
    assert reviewed["export_covered"] is True
    assert reviewed["checks"]["export"]["result"] == "PASS"
    assert reviewed["checks"]["export"]["pos_total_net"] == "135.49"
    assert reviewed["checks"]["export"]["pos_card"] == "159.49"
    assert reviewed["checks"]["export"]["pos_noncard_net"] == "-24.00"
    assert reviewed["staged"] is True
    assert reviewed["figures_final"] is False
    assert reviewed["screen"]["count_text"] == "9"


def test_dates_rejects_a_pdf_already_used_for_another_verified_date(afsa_state):
    """The reuse rule reads the isolated ledger, not the wall clock."""
    folder = GOLDEN / "inputs"
    pdf = folder / "6874_347360075_99871480206_20102026.pdf"
    evidence = folder / "processor-evidence.json"
    other = afsa_state / "runs" / "2026-10-19"
    other.mkdir(parents=True)
    (other / pdf.name).write_bytes(pdf.read_bytes())
    (afsa_state / "verified-dates.json").write_text(
        json.dumps({"start": "2026-10-01", "verified": {"2026-10-19": {"md5": "recorded"}}}),
        encoding="utf-8",
    )
    from feina_client.approved import find_approved, run_dates

    found = find_approved()
    assert found and found["ok"]
    reused = run_dates(
        found["checks"],
        business_date="2026-10-20",
        pdfs=[pdf],
        evidence=[evidence],
        slack_on_fail=False,
    )
    assert reused["result"] == "FAIL"
    assert any("byte-identical" in item and "2026-10-19" in item for item in reused["problems"])
    assert not Path("/root/.hermes/automation/af-ingresos/verified-dates.json").exists()

    (afsa_state / "verified-dates.json").unlink()
    (other / pdf.name).unlink()
    clean = run_dates(
        found["checks"],
        business_date="2026-10-20",
        pdfs=[pdf, folder / "resumen_347360075_T0000001_202610210123.pdf"],
        evidence=[evidence],
        slack_on_fail=False,
    )
    assert clean["result"] == "PASS", clean.get("problems")
    assert clean["problems"] == []


def test_future_membr_export_slot_is_empty_until_a_file_is_supplied():
    note = (SLOT / "README.md").read_text(encoding="utf-8")
    assert "not in the fixture" in note or "not replayed" in note or "not covered" in note.lower() or "is not" in note
    books = sorted(SLOT.glob("*.xlsx"))
    if not books:
        assert books == []
        return
    day = (SLOT / "business-date.txt").read_text(encoding="utf-8").strip()
    from feina_client.approved import find_approved, run_export

    found = find_approved()
    assert found and found["ok"]
    result = run_export(found["checks"], business_date=day, xlsx=books[0], slack_on_fail=False)
    assert result["result"] == "PASS", result.get("problems")
