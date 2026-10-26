"""The AFSA night calls the approved scripts and stops when they do not pass."""

from __future__ import annotations

import hashlib
import json
import shutil
from datetime import date, datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook

from feina_client.reconcile import review_day
from feina_client.runtime import sha256_file

DAY = date(2026, 10, 20)
MADRID = ZoneInfo("Europe/Madrid")
ROOT = Path(__file__).resolve().parents[1]
APPROVED = ROOT / "files" / "skills" / "client-drop" / "afsa-ingresos" / "approved"
GOLDEN = ROOT / "tests" / "fixtures" / "golden" / "2026-10-20"


def _inputs():
    folder = GOLDEN / "inputs"
    return (
        folder / "daily-payload.json",
        [
            folder / "6874_347360075_99871480206_20102026.pdf",
            folder / "resumen_347360075_T0000001_202610210123.pdf",
        ],
        [folder / "processor-evidence.json"],
    )


def _workbook(path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "2026"
    book.create_sheet("2025")
    sheet["A4"] = "KEEP-A4"
    sheet["B4"] = "4 Enero 2026"
    sheet["C4"] = 1
    sheet["H4"] = "KEEP-H4"
    sheet["H293"] = "KEEP-H293"
    sheet["B294"] = "20 Octubre"
    sheet["H294"] = "KEEP-H294"
    book["2025"]["A1"] = "KEEP-2025"
    sheet.column_dimensions["H"].hidden = False
    book.save(path)


def _export_book(path: Path) -> None:
    import datetime as dt

    book = Workbook()
    sheet = book.active
    sheet.title = "Informe"
    meta = book.create_sheet("Informe Meta Data")
    meta.append(["Nombre Informe", "Fecha de", "Fecha hasta", "Fecha descargado"])
    meta.append(
        [
            "Transaccion Punto de Venta",
            dt.datetime(2026, 10, 20, 0, 0, 0),
            dt.datetime(2026, 10, 20, 23, 59, 59),
            dt.datetime(2026, 10, 21, 10, 4, 29),
        ]
    )
    sheet.append(
        [
            "Fecha de la transacción",
            "Estado de transacción",
            "Tipo de pago",
            "Subtotal",
            "Nombre del gimnasio",
            "Número de recibo",
        ]
    )
    sheet.append(
        [dt.datetime(2026, 10, 20, 13, 27, 35), "Pagado", "Tarjeta", 159.49, "Anytime Fitness Sant Adria", "R1"]
    )
    sheet.append(
        [dt.datetime(2026, 10, 20, 18, 0, 0), "Pagado", "Efectivo", -24, "Anytime Fitness Sant Adria", "R2"]
    )
    book.save(path)


def test_vendored_scripts_match_the_recorded_hashes():
    recorded = json.loads((APPROVED / "hashes.json").read_text(encoding="utf-8"))["files"]
    assert recorded["afsa_checks.py"].startswith("1139bac7")
    assert recorded["update_workbook.py"].startswith("3f5437c3")
    for name, digest in recorded.items():
        assert hashlib.sha256((APPROVED / name).read_bytes()).hexdigest() == digest


def test_a_hash_mismatch_does_not_check_the_day(tmp_path, monkeypatch, client):
    copy = tmp_path / "approved"
    shutil.copytree(APPROVED, copy)
    script = copy / "afsa_checks.py"
    script.write_bytes(script.read_bytes() + b"\n")
    monkeypatch.setenv("FEINA_AFSA_APPROVED", str(copy))
    book = client / "AFSA-Ingresos.xlsx"
    _workbook(book)
    before = sha256_file(book)
    reviewed = review_day(
        business_date=DAY,
        membr=None,
        pdfs=[],
        workbook=book,
        slack_on_fail=False,
    )
    assert reviewed["status"] == "CHECKS_UNAVAILABLE"
    assert reviewed["figures_final"] is False
    assert reviewed["tolerance"] is None
    assert sha256_file(book) == before
    assert "afsa_checks.py" in reviewed["mismatched"]


def test_export_failure_stops_without_a_workbook_write(tmp_path, client):
    bad = tmp_path / "pos-export.xlsx"
    Workbook().save(bad)
    book = client / "AFSA-Ingresos.xlsx"
    _workbook(book)
    before = sha256_file(book)
    reviewed = review_day(
        business_date=DAY,
        membr=bad,
        pdfs=[],
        workbook=book,
        slack_on_fail=False,
    )
    assert reviewed["status"] == "CHECK_FAIL"
    assert reviewed["reason"] == "export"
    assert reviewed["figures_final"] is False
    assert reviewed["problems"]
    assert reviewed["staged"] if False else reviewed.get("staged_file") is None
    assert sha256_file(book) == before


def test_missing_payment_pdfs_are_not_a_zero_day_and_harlands_is_not_required(client):
    payload, _pdfs, _evidence = _inputs()
    book = client / "AFSA-Ingresos.xlsx"
    _workbook(book)
    before = sha256_file(book)
    reviewed = review_day(
        business_date=DAY,
        membr=None,
        pdfs=[],
        evidence=[],
        payload_path=payload,
        workbook=book,
        slack_on_fail=False,
    )
    assert reviewed["status"] == "NEEDS_CONFIRMATION"
    assert reviewed["reason"] == "payment_reports_missing"
    assert reviewed["figures_final"] is False
    assert reviewed["export_covered"] is False
    assert sha256_file(book) == before
    from feina_client.proactive import compose_daily

    message = compose_daily("fr", reviewed, when_clean=True)
    assert "définitif" in message
    assert "135.49" not in message


def test_dates_failure_and_a_payload_that_disagrees_with_export_do_not_write(tmp_path, client):
    payload, pdfs, evidence = _inputs()
    book = client / "AFSA-Ingresos.xlsx"
    _workbook(book)
    before = sha256_file(book)
    wrong_day = review_day(
        business_date=date(2026, 10, 21),
        membr=None,
        pdfs=pdfs,
        evidence=evidence,
        payload_path=payload,
        workbook=book,
        slack_on_fail=False,
    )
    assert wrong_day["status"] == "CHECK_FAIL"
    assert wrong_day["reason"] == "dates"
    assert wrong_day["figures_final"] is False
    assert sha256_file(book) == before

    export = tmp_path / "pos-export.xlsx"
    _export_book(export)
    other = tmp_path / "payload.json"
    body = json.loads(payload.read_text(encoding="utf-8"))
    body["pos_card"] = "1.00"
    other.write_text(json.dumps(body), encoding="utf-8")
    disagreed = review_day(
        business_date=DAY,
        membr=export,
        pdfs=pdfs,
        evidence=evidence,
        payload_path=other,
        workbook=book,
        screen_count_text="999 rows",
        screen_total_text="1.00",
        slack_on_fail=False,
    )
    assert disagreed["reason"] == "payload_does_not_match_export"
    assert disagreed["figures_final"] is False
    assert disagreed["screen"]["captured"] is True
    assert sha256_file(book) == before


def test_missing_workbook_is_not_created(client):
    payload, pdfs, evidence = _inputs()
    missing = client / "missing.xlsx"
    reviewed = review_day(
        business_date=DAY,
        membr=None,
        pdfs=pdfs,
        evidence=evidence,
        payload_path=payload,
        workbook=missing,
        slack_on_fail=False,
    )
    assert reviewed["reason"] == "workbook_missing"
    assert reviewed["figures_final"] is False
    assert not missing.exists()


def test_nightly_without_payment_reports_makes_no_model_calls(client, monkeypatch, tmp_path):
    from feina_client.jobs import run_nightly

    dest = tmp_path / "pos-export.xlsx"
    _export_book(dest)

    def fake_download(cfg, **kwargs):
        out = Path(kwargs["out_dir"])
        out.mkdir(parents=True, exist_ok=True)
        saved = out / "pos-export.xlsx"
        saved.write_bytes(dest.read_bytes())
        return {
            "status": "OK",
            "files": [
                {
                    "report": "pos_transactions",
                    "path": str(saved),
                    "sha256": "abc",
                    "bytes": saved.stat().st_size,
                    "screen_count_text": "Rows: 2",
                    "screen_total_text": "Total 135,49",
                }
            ],
        }

    monkeypatch.setattr("feina_client.jobs.run_download", fake_download)
    book = client / "AFSA-Ingresos.xlsx"
    _workbook(book)
    before = sha256_file(book)
    result = run_nightly(
        {"pos_report_id": "pos_transactions"},
        now=datetime(2026, 10, 21, 4, 5, tzinfo=MADRID),
        session=object(),
        inbox=client / "inbox",
        workbook=book,
        membr_dir=tmp_path / "run",
    )
    assert result["figures_final"] is False
    assert result["model_calls"] == 0
    assert result["reason"] == "payment_reports_missing"
    assert sha256_file(book) == before
